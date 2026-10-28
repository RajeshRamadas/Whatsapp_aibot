import csv
import json

from docx import Document

from app.ingest import Block, chunk_blocks, embedding_text, extract, split_long


def test_markdown_keeps_heading_path(tmp_path):
    path = tmp_path / 'policy.md'
    path.write_text('# Policies\n\nIntro text.\n\n## Refunds\n\nWithin 30 days.\n\n## Shipping\n\nFree above Rs 999.\n')
    pieces = chunk_blocks(extract(path, '.md'))
    assert [(p.heading, p.content) for p in pieces] == [
        ('Policies', 'Intro text.'),
        ('Policies > Refunds', 'Within 30 days.'),
        ('Policies > Shipping', 'Free above Rs 999.'),
    ]


def test_chunks_never_cross_sections_or_pages():
    blocks = [Block('a' * 50, 'A'), Block('b' * 50, 'A'), Block('c' * 50, 'B'), Block('d' * 50, None, page=2)]
    pieces = chunk_blocks(blocks, limit=500)
    assert [(p.heading, p.page) for p in pieces] == [('A', None), ('B', None), (None, 2)]
    assert pieces[0].content == 'a' * 50 + '\n\n' + 'b' * 50


def test_long_text_splits_on_sentences_with_overlap():
    sentences = ' '.join(f'Sentence number {i} is here.' for i in range(40))
    pieces = chunk_blocks([Block(sentences)], limit=200)
    assert len(pieces) > 1
    assert all(len(p.content) <= 200 + 2 + 50 for p in pieces)  # limit plus one carried-over paragraph
    assert all(p.content.rstrip().endswith('.') for p in pieces)
    assert split_long('x' * 450, 200) == ['x' * 200, 'x' * 200, 'x' * 50]


def test_spreadsheet_rows_stay_whole_with_column_names(tmp_path):
    path = tmp_path / 'plans.csv'
    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['plan', 'price', 'notes'])
        w.writerows([['Basic', '99', ''], ['Pro', '499', 'priority support']])
    blocks = extract(path, '.csv')
    assert [b.text for b in blocks] == ['plan: Basic; price: 99', 'plan: Pro; price: 499; notes: priority support']
    assert all(b.row for b in blocks)


def test_json_and_docx(tmp_path):
    j = tmp_path / 'faq.json'
    j.write_text(json.dumps([{'q': 'Open on Sunday?', 'a': 'No'}]))
    assert extract(j, '.json')[0].text == 'q: Open on Sunday?; a: No'

    d = Document()
    d.add_heading('Pricing', 1)
    d.add_paragraph('All prices include GST.')
    table = d.add_table(rows=2, cols=2)
    for (r, c), v in {(0, 0): 'Plan', (0, 1): 'Price', (1, 0): 'Basic', (1, 1): '99'}.items():
        table.cell(r, c).text = v
    path = tmp_path / 'pricing.docx'
    d.save(path)
    blocks = extract(path, '.docx')
    assert [(b.heading, b.text) for b in blocks] == [('Pricing', 'All prices include GST.'),
                                                    ('Pricing', 'Plan: Basic; Price: 99')]


def test_embedding_text_carries_context():
    piece = chunk_blocks([Block('Within 30 days.', 'Refunds', page=3)])[0]
    assert embedding_text('policy.pdf', piece) == 'policy.pdf > Refunds (page 3)\nWithin 30 days.'
